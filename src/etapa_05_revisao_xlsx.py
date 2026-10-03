"""Audita e extrai a planilha original sem modificá-la ou aprovar resultados."""
import hashlib
import html
import json
import math
import re
from pathlib import Path
import openpyxl

SHA_ESPERADO = 'ee262a683bae0db283c645a2e85f8261545c3ba96775f602f8af7c0261e336e4'
SOCIO = 'INFORMAÇÕES SÓCIO DEMOGRAFICAS'
ESTRATEGIA = 'ESTRATÉGIA POR FAIXA DE RETORNO'
PIB = 'PIB E PROJEÇÃO'
POTENCIAL = 'POTENCIAL VENDAS 2025 A 2029'
PAPEIS = {
ESTRATEGIA: 'RESULTADO_DERIVADO_A_REFAZER',
PIB: 'MISTO_ENTRADAS_CANDIDATAS_E_PROJECOES',
'DIFERENÇA': 'RESULTADO_DERIVADO_A_REFAZER',
POTENCIAL: 'RESULTADO_DERIVADO_A_REFAZER',
'TABELA DE PRODUTOS': 'CATALOGO_CANDIDATO_CONFERIR_COM_PDF',
'RANKING': 'PREFERENCIA_SEM_EVIDENCIA_IDENTIFICADA',
'RANKING 2': 'RANKING_SEM_EVIDENCIA_IDENTIFICADA',
SOCIO: 'MISTO_SOCIOECONOMICO_E_PREMISSAS_LOGISTICAS',
'DADOS COMPLEMENTARES': 'PARAMETROS_HISTORICOS_INTERPRETACAO_PENDENTE',
'INDICE': 'DEFINICOES_E_PREMISSAS_ORIGEM_PENDENTE',
}


def gravar(p, dados):
    p.write_text(json.dumps(dados, ensure_ascii=False, indent=2, default=str), encoding='utf-8')


def cidade(v):
    return re.sub(r'\s+', ' ', str(v).strip()).upper()


def executar(raiz: Path, run_id: str):
    quality = raiz / 'quality' / '05_revisao_xlsx' / run_id
    quality.mkdir(parents=True, exist_ok=False)
    pasta = raiz / 'datalake' / '01_bronze' / run_id / 'xlsx_original'
    pasta.mkdir(parents=True, exist_ok=False)
    erros, alertas, reconciliacoes = [], [], []
    resumo = {}
    livro = None
    def alerta(codigo, motivo, aba=None, celula=None, **extra):
        alertas.append({'codigo': codigo, 'motivo': motivo, 'aba': aba, 'celula': celula, **extra})
    def comparar(tipo, municipio, observado, calculado, aba, celula, regra):
        ok = math.isclose(observado, calculado, rel_tol=1e-10, abs_tol=1e-6)
        reconciliacoes.append({'tipo': tipo, 'cidade': municipio, 'aba': aba, 'celula': celula, 'valor_original': observado, 'valor_reproduzido': calculado, 'regra_inferida': regra, 'coincide': ok, 'interpretacao': 'REPRODUCAO_ARITMETICA_NAO_VALIDA_METODO'})
    try:
        arquivos = sorted(p for p in (raiz / 'datalake' / '00_landing').rglob('*.xlsx') if not p.name.startswith('~$'))
        if len(arquivos) != 1:
            raise ValueError('Esperado um XLSX original; revisar o perfil para novas fontes')
        arquivo = arquivos[0]
        sha = hashlib.sha256(arquivo.read_bytes()).hexdigest()
        if sha != SHA_ESPERADO:
            raise ValueError('XLSX alterado ou sem perfil validado; revisar layout antes de extrair')
        livro = openpyxl.load_workbook(arquivo, data_only=False)
        if set(livro.sheetnames) != set(PAPEIS):
            raise ValueError('Abas incompatíveis com o perfil')
        origem = {'arquivo_origem': arquivo.relative_to(raiz).as_posix(), 'source_sha256': sha}
        celulas, abas, tabelas = [], [], {}
        for s in livro:
            linhas = []
            formulas = 0
            for row in s:
                for c in row:
                    if c.value is None:
                        continue
                    formulas += c.data_type == 'f'
                    celulas.append({**origem, 'aba': s.title, 'celula': c.coordinate, 'linha': c.row, 'coluna': c.column, 'valor_original': c.value, 'tipo_excel': c.data_type, 'formato_excel': c.number_format})
                if row[0].row > 1 and any(c.value is not None for c in row):
                    linhas.append({**origem, 'aba': s.title, 'linha_origem': row[0].row, 'campos': [{'coluna': c.column_letter, 'cabecalho_original': s.cell(1,c.column).value, 'celula': c.coordinate, 'valor_original': c.value} for c in row]})
            tabelas[s.title] = linhas
            abas.append({'aba': s.title, 'linhas_layout': s.max_row, 'colunas_layout': s.max_column, 'celulas_preenchidas': sum(c.value is not None for row in s for c in row), 'formulas_armazenadas': formulas, 'papel': PAPEIS[s.title], 'situacao': 'EXTRAIDO_PENDENTE_VALIDACAO'})
        gravar(pasta / 'celulas_originais.json', celulas)
        gravar(pasta / 'abas_classificadas.json', abas)
        gravar(pasta / 'linhas_originais.json', tabelas)
        total_formulas = sum(a['formulas_armazenadas'] for a in abas)
        alerta('XLSX001', 'Resultados armazenados como valores. Não há fórmulas recuperáveis nesta versão.' if not total_formulas else 'Fórmulas armazenadas exigem revisão.', formulas_armazenadas=total_formulas)
        socio = livro[SOCIO]
        municipios = [cidade(socio.cell(i,1).value) for i in range(2,14)]
        if len(set(municipios)) != 12:
            raise ValueError('Esperadas 12 cidades distintas na aba socioeconômica')
        indicadores = []
        for i in range(2,14):
            for j in range(2,25):
                c=socio.cell(i,j)
                indicadores.append({**origem, 'cidade_chave': cidade(socio.cell(i,1).value), 'cidade_original': socio.cell(i,1).value, 'aba': SOCIO, 'celula': c.coordinate, 'indicador_original': socio.cell(1,j).value, 'valor_original': c.value, 'papel': 'ENTRADA_SOCIOECONOMICA_CANDIDATA' if j<19 else 'PREMISSA_LOGISTICA_A_REVALIDAR', 'fonte_institucional_informada': ('IBGE informado pelo usuário; tabela específica pendente' if j==3 else 'IBGE/Caravela, atribuição por indicador pendente') if j<19 else None, 'periodo_referencia': None, 'unidade_validada': None, 'status': 'ORIGEM_PERIODO_UNIDADE_PENDENTES'})
        gravar(pasta / 'indicadores_socioeconomicos_originais.json', indicadores)
        alerta('XLSX002', 'Cabeçalho POPULAÇÃO EM MILHARES contrasta com números como 364331. Confirmar unidade sem dividir automaticamente.', SOCIO, 'B1')
        alerta('XLSX003', 'PIB de Frutal aparece como 2473117,87 enquanto outras cidades usam ordens de grandeza muito maiores. Confirmar reais versus milhares de reais na fonte.', SOCIO, 'C4', valor_original=socio['C4'].value)
        for i in range(2,14):
            classes=sum(socio.cell(i,j).value for j in range(14,19))
            if not math.isclose(classes,1,abs_tol=0.00001):
                alerta('XLSX004', 'Classes A–E não somam exatamente 100%. Pode decorrer de arredondamento; conferir definições e fonte.', SOCIO, f'N{i}:R{i}', cidade=municipios[i-2], soma=classes)
        estrategia=livro[ESTRATEGIA]
        for i in range(2,14):
            for j,limite in [(2,90),(3,180)]:
                c=estrategia.cell(i,j)
                if not -limite<=c.value<=limite:
                    alerta('XLSX005', 'Coordenada fora do intervalo em graus decimais. Não inferir posição da vírgula automaticamente.', ESTRATEGIA, c.coordinate, cidade=cidade(estrategia.cell(i,1).value), valor_original=c.value)
            comparar('RAZAO_POTENCIAL_CUSTO',cidade(estrategia.cell(i,1).value),estrategia.cell(i,6).value,estrategia.cell(i,4).value/estrategia.cell(i,5).value,ESTRATEGIA,f'F{i}','potencial / custo_entrega')
        for nome in [PIB,POTENCIAL]:
            s=livro[nome]
            for i in range(2,14):
                base=s.cell(i,2).value;taxa=s.cell(i,3).value
                for j,ano in enumerate(range(2025,2030),4):
                    comparar('PROJECAO_ORIGINAL',cidade(s.cell(i,1).value),s.cell(i,j).value,base*(1+taxa)**(ano-2023),nome,s.cell(i,j).coordinate,'base * (1 + taxa_rotulada_cinco_anos) ** (ano - 2023)')
        por_cidade={cidade(livro[POTENCIAL].cell(i,1).value):i for i in range(2,14)}
        for i in range(2,14):
            s=livro['DIFERENÇA'];municipio=cidade(s.cell(i,1).value);r=por_cidade[municipio];v25=livro[POTENCIAL].cell(r,4).value;v29=livro[POTENCIAL].cell(r,8).value
            comparar('DIFERENCA_ABSOLUTA',municipio,s.cell(i,2).value,v29-v25,'DIFERENÇA',f'B{i}','potencial_2029 - potencial_2025')
            comparar('DIFERENCA_PERCENTUAL',municipio,s.cell(i,3).value,(v29/v25-1)*100,'DIFERENÇA',f'C{i}','(potencial_2029 / potencial_2025 - 1) * 100')
        alerta('XLSX006', 'Projeções reproduzem a taxa rotulada como crescimento em cinco anos a cada ano. Confirmar se ela é acumulada ou anualizada e seus períodos antes do recálculo.', PIB, 'C2:H13')
        alerta('XLSX007', 'As projeções de potencial reproduzem crescimento do PIB sem evidência de relação validada com vendas. Não converter PIB em demanda automaticamente.', POTENCIAL)
        alerta('XLSX008', 'Razão potencial/custo de uma entrega não demonstra lucro nem retorno por real investido: períodos, custos comerciais e volume precisam ser compatíveis.', ESTRATEGIA, 'D1:G13')
        alerta('XLSX009', 'Todos os municípios apontam Ribeirão Preto. Distâncias, endereço de origem e trecho simples versus ida/volta não estão documentados. Comparar fábrica e quatro CDs.', SOCIO, 'S2:T13')
        alerta('XLSX010', 'Percentuais de preferência e rankings não têm evidência específica identificada no XLSX; soma 100% não comprova pesquisa.', 'RANKING', 'D2:D11', soma=sum(livro['RANKING'].cell(i,4).value for i in range(2,12)))
        alerta('XLSX011', 'Percentuais históricos de pessoas que referiram consumo não representam frequência de compra, quantidade de produtos ou vendas capturáveis. Classes A–E não equivalem automaticamente às faixas de renda dos estudos.', 'DADOS COMPLEMENTARES')
        alerta('XLSX012', 'Valor de mercado de 2021 e critérios de classes requerem fonte, abrangência e definição específica.', 'INDICE')
        alerta('XLSX013', 'Cabeçalho de crescimento 2024–2025 não possui coluna de nível de PIB 2024 no arquivo.', PIB, 'I1')
        gravar(pasta / 'reconciliacao_aritmetica_original.json', reconciliacoes)
        gravar(pasta / 'alertas_revisao.json', alertas)
        doc='<!doctype html><meta charset="utf-8"><title>Revisão XLSX original</title><style>table{border-collapse:collapse;margin-bottom:24px}td,th{border:1px solid #aaa;padding:6px;vertical-align:top}</style><h1>Revisão do XLSX original</h1><p>Valores preservados. Reprodução aritmética não aprova metodologia. Fontes, unidades, períodos e premissas permanecem pendentes.</p>'
        for titulo,dados in [('Abas',abas),('Alertas',alertas),('Reconciliação',reconciliacoes)]:
            cols=list(dict.fromkeys(k for r in dados for k in r));doc+='<h2>'+titulo+'</h2><table><tr>'+''.join('<th>'+html.escape(k)+'</th>' for k in cols)+'</tr>'+''.join('<tr>'+''.join('<td>'+html.escape(str(r.get(k,'')))+'</td>' for k in cols)+'</tr>' for r in dados)+'</table>'
        (pasta/'conferencia_xlsx.html').write_text(doc,encoding='utf-8')
        resumo={**origem,'abas':len(abas),'formulas_armazenadas':total_formulas,'celulas_preenchidas':len(celulas),'cidades':municipios,'indicadores_e_premissas_extraidos':len(indicadores),'comparacoes_aritmeticas':len(reconciliacoes),'comparacoes_coincidentes':sum(r['coincide'] for r in reconciliacoes),'comparacoes_divergentes':sum(not r['coincide'] for r in reconciliacoes),'alertas':len(alertas),'saida':pasta.relative_to(raiz).as_posix()}
    except Exception as e:
        erros.append({'erro':str(e)})
    finally:
        if livro is not None:livro.close()
    resultado={'etapa':'05_revisao_xlsx','run_id':run_id,'status':'FALHA' if erros else 'REVISADO_COM_PENDENCIAS_METODOLOGICAS','resumo':resumo,'erros':erros,'recalculo_aprovado':False,'original_modificado':False,'proxima_etapa':'Validar fontes, unidades, períodos, premissas comerciais e logística antes de refazer cálculos.'}
    gravar(quality/'conclusao_05_revisao_xlsx.json',resultado)
    if erros:raise RuntimeError('Falha na revisão do XLSX; consultar conclusão da etapa 05')
    return resultado
