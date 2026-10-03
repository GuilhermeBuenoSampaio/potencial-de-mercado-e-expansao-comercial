"""Coleta municipal oficial e rastreavel; nao altera a analise original."""
import argparse
import gzip
import hashlib
import html
import json
import math
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.parse import quote
from uuid import uuid4

PROJETO = 'potencial-de-mercado-e-expansao-comercial'
API = 'https://servicodados.ibge.gov.br/api'
CIDADES = [('Franca','SP'),('Cristais Paulista','SP'),('Frutal','MG'),('Morro Agudo','SP'),('Sales Oliveira','SP'),('Orlândia','SP'),('Nuporanga','SP'),('Ipuã','SP'),('Guaíra','SP'),('Barretos','SP'),('Planura','MG'),('Colômbia','SP')]
ORIGENS = [('São Carlos','SP','fabrica'),('Ribeirão Preto','SP','cd'),('Campinas','SP','cd'),('Sorocaba','SP','cd'),('São Paulo','SP','cd')]
POLITICA = {'versao':'1.0', 'data':'2026-09-30', 'populacao':1, 'emprego':2, 'pib':3,
 'censo':'ESTRUTURAL: ultimo censo disponivel, sempre com ano de referencia',
 'geografia':'ultimo produto oficial disponivel; versao desconhecida implica conferencia pendente',
 'rotas_dias':90, 'combustivel_dias':30,
 'criterios':['instituicao identificada e fonte primaria preferencial','metodologia e universo explicitados','periodo, unidade e territorio identificados','codigo IBGE municipal validado','consulta reproduzivel e resposta original preservada','limitacoes, revisoes e lacunas registradas'],
 'regra':'Data de coleta nao substitui ano de referencia. Dado antigo nao e apagado nem convertido em atual. Ausencia nao e zero. Limites de defasagem sao escolhas do projeto, nao normas do IBGE.'}

def gravar(p, obj):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')

def chave(s):
    return ''.join(c for c in unicodedata.normalize('NFKD',s) if not unicodedata.combining(c)).upper()

def numero(s):
    try:
        d = Decimal(str(s))
        return float(d) if d.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None

class Cliente:
    def __init__(self, pasta):
        self.pasta=pasta; self.fontes=[]; self.falhas=[]; self.cache={}
    def obter(self, url):
        if url in self.cache: return self.cache[url]
        erro=None
        for tentativa in range(3):
            try:
                with urlopen(Request(url,headers={'User-Agent':PROJETO+'/1.0','Accept':'application/json'}),timeout=45) as r:
                    bruto=r.read(); final=r.url; status=r.status
                corpo=gzip.decompress(bruto) if bruto.startswith(b'\x1f\x8b') else bruto
                dados=json.loads(corpo.decode('utf-8-sig'))
                sha=hashlib.sha256(corpo).hexdigest(); nome=hashlib.sha256(url.encode()).hexdigest()[:20]+'.json'
                self.pasta.mkdir(parents=True,exist_ok=True); (self.pasta/nome).write_bytes(corpo)
                fonte={'url':url,'url_final':final,'http_status':status,'coleta_utc':datetime.now(timezone.utc).isoformat(),'arquivo':str(self.pasta/nome),'sha256':sha,'sha256_transporte':hashlib.sha256(bruto).hexdigest(),'instituicao':'IBGE'}
                self.fontes.append(fonte); self.cache[url]=(dados,fonte)
                return dados,fonte
            except Exception as e:
                erro=f'{type(e).__name__}: {e}'
                if tentativa<2: time.sleep(tentativa+1)
        self.falhas.append({'url':url,'erro':erro})
        raise RuntimeError(erro)

def classificacoes(meta, especiais=None):
    especiais=especiais or {}; partes=[]
    for c in meta.get('classificacoes',[]):
        cid=str(c['id'])
        ids=especiais.get(cid)
        if ids is None:
            totais=[str(x['id']) for x in c['categorias'] if x['nome']=='Total']
            if len(totais)!=1: raise ValueError(f'Total ambiguo na classificacao {cid}')
            ids=totais
        validos={str(x['id']) for x in c['categorias']}
        if not set(ids)<=validos: raise ValueError(f'Categoria desconhecida: {cid}')
        partes.append(cid+'['+','.join(ids)+']')
    return '|'.join(partes)

def distancia(a,b):
    lat1,lon1,lat2,lon2=map(math.radians,(a['latitude'],a['longitude'],b['latitude'],b['longitude']))
    v=math.sin((lat2-lat1)/2)**2+math.cos(lat1)*math.cos(lat2)*math.sin((lon2-lon1)/2)**2
    return round(6371.0088*2*math.asin(math.sqrt(min(1,max(0,v)))),3)

def executar(raiz, run_id, comparar_original=False):
    raiz=Path(raiz).resolve()
    qualidade=raiz/'quality'/'06_coleta_municipal'/run_id
    qualidade.mkdir(parents=True,exist_ok=False)
    bronze=raiz/'datalake'/'01_bronze'/run_id/'coleta_municipal'
    cliente=Cliente(raiz/'datalake'/'00_landing'/'fontes_externas'/run_id/'ibge')
    observacoes=[]; municipios=[]; pendencias=[]; ano=datetime.now(timezone.utc).year
    try:
        cad={}
        for uf in ['SP','MG']:
            lista,_=cliente.obter(API+'/v1/localidades/estados/'+uf+'/municipios')
            for m in lista: cad[(chave(m['nome']),uf)]=m
        for nome,uf in CIDADES+[(n,u) for n,u,_ in ORIGENS]:
            m=cad.get((chave(nome),uf))
            if not m: raise ValueError(f'Municipio nao identificado: {nome}/{uf}')
            municipios.append({'codigo_ibge':str(m['id']),'municipio':m['nome'],'uf':uf,'alvo':(nome,uf) in CIDADES,'regionalizacao':m})
        ids=','.join(m['codigo_ibge'] for m in municipios if m['alvo'])
        por_id={m['codigo_ibge']:m for m in municipios}
        # Um territorio municipal unico; nenhuma agregacao estadual substitui municipio.
        specs=[('6579',['9324'],'populacao',None),('5938',['37','516','520','6574','528'],'pib',None),
               ('9509',['706','707','708','1606','10143','662'],'emprego',None),
               ('4714',['93','6318','614'],'censo',None),('10295',['13431','13534'],'censo',None),
               ('10296',['13604','1013604'],'censo',{'386':['9680','9681','9682','9683','9684','9685','9686','9687','9688','9689','9690','9692']}),
               ('9923',['93','1000093'],'censo',{'1':['6795','1','2']}),
               ('9528',['706','707'],'emprego',{'12762':['117897','117549','117443','116944']})]
        for tabela,vars_,grupo,especial in specs:
            print('Consultando tabela '+tabela, flush=True)
            try:
                meta,mfonte=cliente.obter(API+'/v3/agregados/'+tabela+'/metadados')
                if 'N6' not in meta['nivelTerritorial'].get('Administrativo',[]):
                    raise ValueError('Tabela sem nivel municipal N6')
                existentes={str(v['id']) for v in meta['variaveis']}
                if not set(vars_)<=existentes: raise ValueError('Variavel ausente na tabela')
                periodos,_=cliente.obter(API+'/v3/agregados/'+tabela+'/periodos')
                anos=sorted(int(p['id']) for p in periodos if str(p['id']).isdigit() and len(str(p['id']))==4 and int(p['id'])<=ano)
                if not anos: raise ValueError('Sem periodos anuais identificaveis')
                selecionados=anos[-6:] if grupo in ('pib','populacao') else anos[-3:] if grupo=='emprego' else anos[-1:]
                url=API+'/v3/agregados/'+tabela+'/periodos/'+'|'.join(map(str,selecionados))+'/variaveis/'+'|'.join(vars_)+'?localidades=N6['+ids+']'
                filtro=classificacoes(meta,especial)
                if filtro: url+='&classificacao='+filtro
                dados,fonte=cliente.obter(url)
                for variavel in dados:
                    for resultado in variavel['resultados']:
                        cats=resultado.get('classificacoes',[])
                        for serie in resultado['series']:
                            codigo=str(serie['localidade']['id'])
                            if codigo not in por_id or not por_id[codigo]['alvo']: raise ValueError('Territorio inesperado na resposta')
                            for periodo,valor in serie['serie'].items():
                                n=numero(valor); unidade=variavel['unidade']; mult=1000 if unidade.lower()=='mil reais' else 1
                                idade=ano-int(periodo)
                                status='ESTRUTURAL' if grupo=='censo' else ('DENTRO_DO_LIMITE' if idade<=POLITICA[grupo] else 'DEFASADO')
                                observacoes.append({'codigo_ibge':codigo,'municipio':por_id[codigo]['municipio'],'uf':por_id[codigo]['uf'],
                                  'tabela':tabela,'variavel_id':str(variavel['id']),'indicador':variavel['variavel'],'classificacoes':cats,
                                  'ano_referencia':int(periodo),'data_publicacao_ou_revisao':next((p.get('modificacao') for p in periodos if str(p['id'])==str(periodo)),None),'ultimo_periodo_tabela':max(anos),'valor_original':valor,'unidade_original':unidade,
                                  'valor':None if n is None else n*mult,'unidade':'R$' if mult==1000 else unidade,'multiplicador':mult,
                                  'grupo':grupo,'status_valor':'NUMERICO' if n is not None else 'MARCADOR_PRESERVADO',
                                  'status_atualidade':status,'defasagem_anos':idade,'url':url,'sha256':fonte['sha256'],
                                  'arquivo_fonte':fonte['arquivo'],'coleta_utc':fonte['coleta_utc'],'url_metadados':mfonte['url']})
            except Exception as e:
                pendencias.append({'item':'tabela_'+tabela,'motivo':str(e)})
        # PIB por habitante oficial: nao dividir pelo numero de moradores de outro ano.
        try:
            meta, mf = cliente.obter(API+'/v1/pesquisas/38/indicadores')
            def localizar(lista):
                for item in lista:
                    if str(item.get('id'))=='47001': return item
                    encontrado=localizar(item.get('children',[]))
                    if encontrado: return encontrado
                return None
            indicador_meta=localizar(meta)
            if not indicador_meta or indicador_meta.get('unidade',{}).get('id')!='R$' or indicador_meta['unidade'].get('multiplicador')!=1:
                raise ValueError('Unidade/metadados do PIB per capita divergentes; conferir antes de normalizar')
            def pib_pc(m):
                try:
                    url=API+'/v1/pesquisas/38/indicadores/47001/resultados/'+m['codigo_ibge']
                    dados,f=cliente.obter(url)
                    for item in dados:
                        if str(item['id'])!='47001': continue
                        for serie in item['res']:
                            if str(serie['localidade']) not in (m['codigo_ibge'],m['codigo_ibge'][:6]):
                                raise ValueError('Territorio inesperado na pesquisa de PIB per capita')
                            for periodo,v in serie['res'].items():
                                if not str(periodo).isdigit() or int(periodo)>ano: continue
                                n=numero(v)
                                if n is None: continue
                                observacoes.append({'codigo_ibge':m['codigo_ibge'],'municipio':m['municipio'],'uf':m['uf'],
                                  'tabela':'pesquisa_38','variavel_id':'47001','indicador':'PIB per capita, serie revisada',
                                  'classificacoes':[],'ano_referencia':int(periodo),'valor_original':v,'unidade_original':'R$',
                                  'valor':n,'unidade':'R$','multiplicador':1,'grupo':'pib','status_valor':'NUMERICO',
                                  'status_atualidade':'DENTRO_DO_LIMITE' if ano-int(periodo)<=POLITICA['pib'] else 'DEFASADO',
                                  'defasagem_anos':ano-int(periodo),'url':url,'sha256':f['sha256'],'arquivo_fonte':f['arquivo'],
                                  'coleta_utc':f['coleta_utc'],'url_metadados':mf['url'],
                                  'notas_fonte':serie.get('notas',{}),'limitacao':'Producao por habitante; nao e renda domiciliar mensal'})
                except Exception as e: return {'item':'PIBpc_'+m['codigo_ibge'],'motivo':str(e)}
            with ThreadPoolExecutor(max_workers=3) as pool:
                for falha in pool.map(pib_pc,[m for m in municipios if m['alvo']]):
                    if falha: pendencias.append(falha)
        except Exception as e:
            pendencias.append({'item':'PIBpc_metadados','motivo':str(e)})
        # Centroide municipal: nao representa endereco da fabrica/CD nem percurso rodoviario.
        def geografia(m):
            try:
                dados,fonte=cliente.obter(API+'/v3/malhas/municipios/'+m['codigo_ibge']+'/metadados')
                d=next(x for x in dados if str(x['id'])==m['codigo_ibge'])
                lat=float(d['centroide']['latitude']);lon=float(d['centroide']['longitude'])
                if not (-90<=lat<=90 and -180<=lon<=180): raise ValueError('Coordenadas invalidas')
                m.update({'latitude':lat,'longitude':lon,'tipo_coordenada':'CENTROIDE_MUNICIPAL','area_metadados':d.get('area'),
                          'url_geografia':fonte['url'],'sha256_geografia':fonte['sha256'],'versao_malha':'NAO_EXPLICITADA_PELO_ENDPOINT'})
            except Exception as e: return {'item':'geografia_'+m['codigo_ibge'],'motivo':str(e)}
        with ThreadPoolExecutor(max_workers=3) as pool:
            for falha in pool.map(geografia,municipios):
                if falha: pendencias.append(falha)
        rotas=[]
        for alvo in [m for m in municipios if m['alvo'] and 'latitude' in m]:
            for origem in [m for m in municipios if not m['alvo'] and 'latitude' in m]:
                rotas.append({'destino':alvo['municipio'],'destino_codigo':alvo['codigo_ibge'],'origem':origem['municipio'],'origem_codigo':origem['codigo_ibge'],
                  'distancia_geodesica_centroides_km':distancia(alvo,origem),'metodo':'Haversine, raio medio 6371.0088 km',
                  'uso':'Comparacao geografica preliminar; nao usar para custo, tempo ou escolha definitiva do CD',
                  'fontes_sha256':[alvo['sha256_geografia'],origem['sha256_geografia']]})
        derivados=[]
        for m in [x for x in municipios if x['alvo']]:
            obs=[x for x in observacoes if x['codigo_ibge']==m['codigo_ibge'] and x['valor'] is not None]
            pibs={x['ano_referencia']:x for x in obs if x['tabela']=='5938' and x['variavel_id']=='37'}
            if pibs:
                fim=max(pibs);inicio=fim-5
                if inicio in pibs and pibs[inicio]['valor']>0 and pibs[fim]['valor']>0:
                    r=pibs[fim]['valor']/pibs[inicio]['valor']
                    derivados.append({'codigo_ibge':m['codigo_ibge'],'municipio':m['municipio'],'indicador':'crescimento_PIB_nominal_5_anos',
                      'ano_inicial':inicio,'ano_final':fim,'acumulado_percentual':(r-1)*100,'taxa_anualizada_percentual':(r**(1/5)-1)*100,
                      'formula':'acumulado=(PIBfinal/PIBinicial-1)*100; anualizada=((PIBfinal/PIBinicial)^(1/5)-1)*100',
                      'fontes_sha256':[pibs[inicio]['sha256'],pibs[fim]['sha256']], 'limitacao':'Precos correntes; crescimento nominal, sem correcao inflacionaria. Nao e previsao de vendas.'})
                else: pendencias.append({'item':'crescimento_5_anos_'+m['codigo_ibge'],'motivo':'Sem par exato de anos separado por cinco anos'})
        # A comparacao preserva celulas e nomes. Sem periodo/unidade originais verificados, nenhuma diferenca e aprovada automaticamente.
        comparacao=[]
        if comparar_original:
            try:
                from openpyxl import load_workbook
                arquivos=[p for p in (raiz/'datalake'/'00_landing').glob('*.xlsx') if not p.name.startswith('~$')]
                if len(arquivos)!=1: raise ValueError('Esperado um XLSX original diretamente em 00_landing')
                arquivo=arquivos[0]; sha=hashlib.sha256(arquivo.read_bytes()).hexdigest(); wb=load_workbook(arquivo,data_only=True,read_only=True)
                try:
                    ws=wb['INFORMAÇÕES SÓCIO DEMOGRAFICAS']; linhas=list(ws.iter_rows())
                    mapa={2:('6579','9324'),3:('5938','37'),6:('9509','1606'),7:('9509','707')}
                    for linha in linhas[1:]:
                        nome=linha[0].value
                        if not nome: continue
                        matches=[m for m in municipios if m['alvo'] and chave(m['municipio'])==chave(str(nome).split(' - ')[0])]
                        for col,(tab,var) in mapa.items():
                            candidatos=[o for o in observacoes if len(matches)==1 and o['codigo_ibge']==matches[0]['codigo_ibge'] and o['tabela']==tab and o['variavel_id']==var and o['valor'] is not None]
                            novo=max(candidatos,key=lambda x:x['ano_referencia']) if candidatos else None
                            comparacao.append({'cidade_original':nome,'campo_original':linhas[0][col-1].value,'celula':linha[col-1].coordinate,
                               'valor_original':linha[col-1].value,'formato_original':linha[col-1].number_format,'arquivo_original':str(arquivo),'sha256_original':sha,
                               'observacao_atual':novo,'status':'COMPARABILIDADE_PENDENTE: validar periodo, unidade e definicao originais'})
                finally: wb.close()
            except Exception as e: pendencias.append({'item':'comparacao_original','motivo':str(e)})
        pendencias.extend([
          {'item':'RENDA PER CAPITA original','motivo':'Definicao ambigua; PIB por habitante e renda domiciliar mensal sao distintos. Foram coletadas rendas media e mediana do Censo. Nao substituir automaticamente.'},
          {'item':'Classes A B C D E','motivo':'Faixas originais sem metodologia suficiente. Dados oficiais de renda per capita nao reproduzem automaticamente classes de renda familiar; intervalo 3–5 SM nao permite separar limiar 4 SM.'},
          {'item':'Servicos industria comercio agricultura','motivo':'Confirmar se as proporcoes originais descrevem PIB, emprego ou estabelecimentos. VAB setorial nao separa comercio e servicos; periodos 2022/2023 sem detalhamento na tabela 5938.'},
          {'item':'percentual populacao ocupada','motivo':'Pessoal ocupado do CEMPRE refere-se a postos nas organizacoes locais; nao equivale ao percentual de residentes empregados.'},
          {'item':'CD mais proximo e quilometros','motivo':'Exige enderecos reais, percurso rodoviario e criterios de atendimento. Distancia entre centroides nao responde definitivamente.'},
          {'item':'quatro consumos Fiorino','motivo':'Exige ano, motorizacao e versao do veiculo, fonte PBE/INMETRO e validacao do consumo observado.'},
          {'item':'ranking produtos consumo e mercado R$81 bilhoes','motivo':'Necessaria fonte, universo, periodo e metodologia verificaveis. Pesquisa historica nao e previsao atual de venda.'}])
        gravar(bronze/'indicadores_municipais.json',observacoes); gravar(bronze/'municipios_geografia.json',municipios)
        gravar(bronze/'distancias_geodesicas.json',rotas);gravar(bronze/'indices_recalculados.json',derivados)
        gravar(bronze/'comparacao_original.json',comparacao);gravar(bronze/'pendencias.json',pendencias)
        numericos=[o for o in observacoes if o['valor'] is not None]
        cobertura=[]
        for tabela,variaveis,_,_ in specs:
            for var in variaveis:
                presentes={o['codigo_ibge'] for o in numericos if o['tabela']==tabela and o['variavel_id']==var}
                faltantes=[m['codigo_ibge'] for m in municipios if m['alvo'] and m['codigo_ibge'] not in presentes]
                cobertura.append({'tabela':tabela,'variavel_id':var,'municipios_com_valor':len(presentes),'municipios_sem_valor':faltantes})
                if faltantes: pendencias.append({'item':'cobertura_'+tabela+'_'+var,'motivo':'Municipios sem valor numerico','codigos':faltantes})
        gravar(bronze/'cobertura_indicadores.json',cobertura)
        gravar(bronze/'pendencias.json',pendencias)
        # Ultimo valor numerico por municipio/variavel/categoria; anos antigos continuam disponiveis no arquivo completo.
        ultimos={}
        for o in numericos:
            k=(o['codigo_ibge'],o['tabela'],o['variavel_id'],json.dumps(o['classificacoes'],sort_keys=True))
            if k not in ultimos or o['ano_referencia']>ultimos[k]['ano_referencia']: ultimos[k]=o
        gravar(bronze/'indicadores_ultimo_disponivel.json',list(ultimos.values()))
        linhas=''.join('<tr>'+''.join('<td>'+html.escape(str(o.get(k,'')))+'</td>' for k in ['municipio','tabela','indicador','classificacoes','valor','unidade','ano_referencia','status_atualidade'])+'</tr>' for o in ultimos.values())
        bronze.mkdir(parents=True,exist_ok=True)
        (bronze/'consulta_indicadores.html').write_text('<!doctype html><html lang="pt-BR"><meta charset="utf-8"><title>Coleta municipal</title><style>body{font:14px Arial;margin:24px}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:8px}th{position:sticky;top:0;background:#eee}</style><h1>'+PROJETO+'</h1><p>Ultimo valor numerico disponivel por indicador e categoria. Ano de referencia e defasagem devem ser conferidos. Classificacoes diferentes nao devem ser somadas indiscriminadamente.</p><table><thead><tr><th>Municipio</th><th>Tabela</th><th>Indicador</th><th>Classificacao</th><th>Valor</th><th>Unidade</th><th>Ano</th><th>Atualidade</th></tr></thead><tbody>'+linhas+'</tbody></table></html>',encoding='utf-8')
        resultado={'projeto':PROJETO,'run_id':run_id,'status':'COLETA_PARCIAL_COM_PENDENCIAS' if cliente.falhas or any(p['item'].startswith(('tabela_','geografia_','PIBpc_','cobertura_')) for p in pendencias) else 'COLETA_CONCLUIDA_COM_CONFERENCIA_PENDENTE',
          'municipios_alvo':sum(m['alvo'] for m in municipios),'observacoes':len(observacoes),'observacoes_numericas':len(numericos),'fontes':len(cliente.fontes),'distancias_geodesicas':len(rotas),'pendencias':pendencias,'falhas_http':cliente.falhas,'politica':POLITICA,'saida':str(bronze)}
        if not numericos: resultado['status']='FALHA_SEM_DADOS'
        gravar(qualidade/'conclusao_06_coleta_municipal.json',resultado)
        if not numericos: raise RuntimeError('Nenhum dado numerico coletado; consulte conclusao e fontes')
        return resultado
    except Exception as e:
        gravar(qualidade/'falha.json',{'erro':str(e),'run_id':run_id,'falhas_http':cliente.falhas})
        raise
    finally:
        gravar(cliente.pasta/'fontes_consultadas.json',cliente.fontes)
        gravar(qualidade/'politica_fontes_atualidade.json',POLITICA)

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--raiz',type=Path,default=Path(__file__).resolve().parent.parent)
    args=p.parse_args();run_id=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'_'+uuid4().hex[:8]
    r=executar(args.raiz,run_id)
    print(json.dumps({'run_id':run_id,'status':r['status'],'observacoes':r['observacoes'],'saida':r['saida']},ensure_ascii=False,indent=2))
if __name__=='__main__': main()
