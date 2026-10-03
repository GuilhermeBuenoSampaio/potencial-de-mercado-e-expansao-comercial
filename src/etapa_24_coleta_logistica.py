"""Coleta ANP e matriz OSRM. Rotas preliminares entre centroides, sem tarifa inferida."""
import argparse
import copy
import hashlib
import io
import json
import math
import re
import unicodedata
from datetime import date, datetime, timezone, timedelta
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from uuid import uuid4

PROJETO = 'potencial-de-mercado-e-expansao-comercial'
ANP = 'https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas'
OSRM = 'https://router.project-osrm.org'


def normalizar(x):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(x)) if not unicodedata.combining(c)).upper().strip()


def baixar(url, destino, registros):
    req = Request(url, headers={'User-Agent':'potencial-de-mercado-e-expansao-comercial/24'})
    with urlopen(req, timeout=35) as resposta:
        dados = resposta.read(25*1024*1024+1)
        final = resposta.url
        if len(dados)>25*1024*1024:
            raise ValueError('Resposta superior ao limite de 25 MiB')
    destino.write_bytes(dados)
    registros.append({'url_solicitada':url,'url_final':final,'arquivo':destino.name,'sha256':hashlib.sha256(dados).hexdigest(),'bytes':len(dados),'consulta_utc':datetime.now(timezone.utc).isoformat()})
    return dados


def url_planilha(html, hoje):
    urls = re.findall(r'href=[\"\']([^\"\']*resumo_semanal_lpc_\d{4}-\d{2}-\d{2}_\d{4}-\d{2}-\d{2}\.xlsx)[\"\']',html)
    candidatos=[]
    for u in set(urls):
        if urlparse(u).hostname!='www.gov.br':
            continue
        datas=re.findall(r'\d{4}-\d{2}-\d{2}',u)
        inicio,fim=map(date.fromisoformat,datas[-2:])
        if inicio<=hoje and fim>=inicio and fim<=hoje+timedelta(days=6) and (hoje-fim).days<=7:
            candidatos.append((fim,u))
    if not candidatos:
        raise ValueError('ANP: nenhuma planilha semanal recente encontrada (maximo 7 dias apos fim)')
    return max(candidatos)[1]


def extrair_gasolina(dados, hoje):
    import openpyxl
    w=openpyxl.load_workbook(io.BytesIO(dados),read_only=True,data_only=True)
    try:
        if 'MUNICIPIOS' not in w.sheetnames:
            raise ValueError('ANP: aba MUNICIPIOS ausente')
        cab=None;achados=[]
        for numero,row in enumerate(w['MUNICIPIOS'].iter_rows(values_only=True),1):
            norm=[normalizar(v) for v in row]
            if 'MUNICIPIO' in norm and 'PRECO MEDIO REVENDA' in norm:
                cab={v:i for i,v in enumerate(norm)};continue
            if cab is None:
                continue
            def v(chave):return row[cab[chave]]
            if normalizar(v('MUNICIPIO'))=='RIBEIRAO PRETO' and normalizar(v('ESTADO'))=='SAO PAULO' and normalizar(v('PRODUTO'))=='GASOLINA COMUM':
                inicio,fim=v('DATA INICIAL').date(),v('DATA FINAL').date()
                preco=float(v('PRECO MEDIO REVENDA'))
                if inicio>hoje or fim<inicio or (hoje-fim).days>7 or not math.isfinite(preco) or preco<=0 or normalizar(v('UNIDADE DE MEDIDA'))!='R$/L':
                    raise ValueError('ANP: periodo, preco ou unidade invalidos')
                achados.append({'municipio':'Ribeirão Preto','uf':'SP','produto':'GASOLINA COMUM','inicio':inicio.isoformat(),'fim':fim.isoformat(),'preco_medio_brl_l':preco,'preco_minimo_brl_l':float(v('PRECO MINIMO REVENDA')),'preco_maximo_brl_l':float(v('PRECO MAXIMO REVENDA')),'postos_pesquisados':int(v('NUMERO DE POSTOS PESQUISADOS')),'aba':'MUNICIPIOS','linha_excel':numero,'status':'FONTE_PRIMARIA_SEMANAL'})
        if len(achados)!=1:
            raise ValueError(f'ANP: esperado um registro, encontrados {len(achados)}')
        return achados
    finally:w.close()


def calcular_rotas(rotas, pontos, matriz):
    n=len(pontos)
    if matriz.get('code')!='Ok' or len(matriz.get('sources',[]))!=n or len(matriz.get('destinations',[]))!=n:
        raise ValueError('Resposta OSRM ou numero de pontos divergente')
    if matriz.get('fallback_speed_cells'):
        raise ValueError('Nao admitir fallback por linha reta')
    for campo in ('distances','durations'):
        a=matriz[campo]
        if len(a)!=n or any(len(l)!=n for l in a):raise ValueError('Matriz nao quadrada')
        for i,l in enumerate(a):
            for j,v in enumerate(l):
                if v is None or not math.isfinite(v) or v<0 or (i!=j and v<=0):raise ValueError('Matriz contem valor invalido/rota ausente')
    indices={x['municipio']:i for i,x in enumerate(pontos)}
    if len(indices)!=n:raise ValueError('Nomes de municipios duplicados')
    snaps=[]
    for i,p in enumerate(pontos):
        desloc=max(matriz['sources'][i]['distance'],matriz['destinations'][i]['distance'])
        snaps.append({'municipio':p['municipio'],'codigo_ibge':p['codigo_ibge'],'deslocamento_ate_malha_m':desloc,'revisar_ponto':desloc>500,'longitude_original':p['longitude_centroide'],'latitude_original':p['latitude_centroide'],'longitude_roteada':matriz['sources'][i]['location'][0],'latitude_roteada':matriz['sources'][i]['location'][1]})
    resultados=[];trechos=[]
    for r in rotas:
        seq=r['sequencia'].split(' -> ')
        if seq[0]!='Ribeirão Preto' or seq[-1]!='Ribeirão Preto':raise ValueError('Circuito deve retornar a Ribeirao Preto')
        km=segundos=0
        for ordem,(u,v) in enumerate(zip(seq,seq[1:]),1):
            i,j=indices[u],indices[v];d=matriz['distances'][i][j]/1000;t=matriz['durations'][i][j]
            km+=d;segundos+=t
            trechos.append({'circuito_id':r['circuito_id'],'ordem':ordem,'origem':u,'destino':v,'distancia_km':d,'duracao_estimada_min':t/60})
        resultados.append({'circuito_id':r['circuito_id'],'sequencia':r['sequencia'],'distancia_rodoviaria_preliminar_km':km,'tempo_modelado_deslocamento_h':segundos/3600,'numero_cidades':r['numero_cidades'],'unidades_varejo_2024':r['unidades_varejo_2024'],'tipo_distancia':'OSRM_ENTRE_CENTROIDES_AJUSTADOS_A_MALHA','status':'PRELIMINAR_ENDERECOS_PENDENTES','revisar_pontos':any(snaps[indices[x]]['revisar_ponto'] for x in set(seq)),'viabilidade_diaria_confirmada':False,'pedagios_total_brl':None})
    return resultados,trechos,snaps


def verificar_manifesto(pasta, nomes):
    m=json.loads((pasta/'manifesto.json').read_text(encoding='utf-8-sig'))
    hashes={x['arquivo']:x['sha256'] for x in m['exportacoes']}
    if m.get('projeto')!=PROJETO:raise ValueError('Projeto da etapa 22 divergente')
    for nome in nomes:
        if hashes.get(nome)!=hashlib.sha256((pasta/nome).read_bytes()).hexdigest():raise ValueError('Hash da etapa 22 divergente: '+nome)
    return m


def executar(raiz,run_id,pasta22=None,parametros=None,cache=None):
    raiz=Path(raiz)
    if pasta22 is None:
        pastas=sorted((raiz/'datalake/03_gold').glob('*/prioridades_circuitos/manifesto.json'))
        if not pastas:raise ValueError('Etapa 22 nao encontrada')
        pasta22=pastas[-1].parent
    pasta22=Path(pasta22)
    verificar_manifesto(pasta22,('circuitos.json','coordenadas_fontes.json'))
    rotas=json.loads((pasta22/'circuitos.json').read_text(encoding='utf-8-sig'))
    pontos=json.loads((pasta22/'coordenadas_fontes.json').read_text(encoding='utf-8-sig'))
    # Somente as cidades efetivamente presentes nas sequencias; nenhuma geocodificacao oculta.
    nomes={x for r in rotas for x in r['sequencia'].split(' -> ')}
    pontos=[p for p in pontos if p['municipio'] in nomes]
    if len(pontos)!=len(nomes):raise ValueError('Coordenadas faltantes para as sequencias')
    parametros=Path(parametros) if parametros else raiz/'docs/23_custos_carga_maxima/PARAMETROS_CARGA_MAXIMA.json'
    config=json.loads(parametros.read_text(encoding='utf-8-sig'))
    if config.get('projeto')!=PROJETO or set(config['circuitos'])!={r['circuito_id'] for r in rotas}:raise ValueError('Configuracao divergente dos circuitos')
    raw=raiz/'datalake/01_bronze'/run_id/'logistica';out=raiz/'datalake/03_gold'/run_id/'logistica';q=raiz/'quality/24_coleta_logistica'/run_id
    for pasta in (raw,out,q):pasta.mkdir(parents=True,exist_ok=False)
    registros=[];erros=[];comb=[];rod=[];trechos=[];snaps=[];hoje=datetime.now(timezone(timedelta(hours=-3))).date()
    conclusao={'projeto':PROJETO,'run_id':run_id,'inicio_utc':datetime.now(timezone.utc).isoformat(),'entrada_22':str(pasta22),'entradas_sha256':{nome:hashlib.sha256((pasta22/nome).read_bytes()).hexdigest() for nome in ('manifesto.json','circuitos.json','coordenadas_fontes.json')},'parametros_entrada_sha256':hashlib.sha256(parametros.read_bytes()).hexdigest(),'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'modo':'REPROCESSAMENTO_CACHE' if cache else 'COLETA_ONLINE'}
    try:
        try:
            if cache:
                dados=(Path(cache)/'anp_semana.xlsx').read_bytes();(raw/'anp_semana.xlsx').write_bytes(dados)
                url='CACHE_LOCAL_ANP';conclusao['origem_cache']=str(cache)
            else:
                html=baixar(ANP,raw/'anp_pagina.html',registros).decode('utf-8')
                url=url_planilha(html,hoje)
                dados=baixar(url,raw/'anp_semana.xlsx',registros)
            comb=extrair_gasolina(dados,hoje)
            comb[0].update(fonte_url=url,fonte_sha256=hashlib.sha256(dados).hexdigest())
        except Exception as e:erros.append({'etapa':'ANP','erro':str(e)})
        try:
            url_osrm=OSRM+'/table/v1/driving/'+';'.join(f"{p['longitude_centroide']},{p['latitude_centroide']}" for p in pontos)+'?annotations=distance,duration'
            if cache:
                dados=(Path(cache)/'osrm_matriz.json').read_bytes();(raw/'osrm_matriz.json').write_bytes(dados)
            else:dados=baixar(url_osrm,raw/'osrm_matriz.json',registros)
            matriz=json.loads(dados);rod,trechos,snaps=calcular_rotas(rotas,pontos,matriz)
            for v in rod:v.update(fonte_url=url_osrm,fonte_sha256=hashlib.sha256(dados).hexdigest(),data_version_malha=matriz.get('data_version'))
        except Exception as e:erros.append({'etapa':'OSRM','erro':str(e)})
        config=copy.deepcopy(config)
        for v in config['circuitos'].values():
            if comb and not cache:v.update(preco_combustivel_brl_l=comb[0]['preco_medio_brl_l'],fonte_combustivel=comb[0]['fonte_url'],periodo_combustivel=comb[0]['inicio']+' a '+comb[0]['fim'])
        # Distancia preliminar fica em campo separado: liberar para custo total exige conferencia dos pontos.
        combustivel=[]
        for r in rod:
            for c in config['consumos']:
                k=float(c['km_l'])
                if not math.isfinite(k) or k<=0:raise ValueError('Consumo invalido')
                combustivel.append({'circuito_id':r['circuito_id'],'consumo_km_l':k,'natureza_consumo':c['natureza'],'distancia_rodoviaria_preliminar_km':r['distancia_rodoviaria_preliminar_km'],'gasolina_brl_l':comb[0]['preco_medio_brl_l'] if comb else None,'custo_somente_gasolina_preliminar_brl':r['distancia_rodoviaria_preliminar_km']/k*comb[0]['preco_medio_brl_l'] if comb else None,'custo_total_rota_brl':None,'faturamento_equilibrio_brl':None,'status':'PRELIMINAR_COM_PEDAGIOS_DESGASTE_MARGEM_PENDENTES'})
        tabelas={'combustivel_oficial':comb,'circuitos_rodoviarios_preliminares':rod,'trechos_rodoviarios_preliminares':trechos,'ajustes_coordenadas':snaps,'cenarios_somente_gasolina':combustivel}
        import pyarrow as pa
        import pyarrow.parquet as pq
        for nome,linhas in tabelas.items():
            (out/(nome+'.json')).write_text(json.dumps(linhas,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
            if linhas:
                pq.write_table(pa.Table.from_pylist(linhas),out/(nome+'.parquet'))
                if pq.read_table(out/(nome+'.parquet')).to_pylist()!=linhas:raise ValueError('Exportacao divergente: '+nome)
        (out/'PARAMETROS_23_COM_COMBUSTIVEL.json').write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8')
        (raw/'registro_coleta.json').write_text(json.dumps({'fontes':registros,'erros':erros},ensure_ascii=False,indent=2),encoding='utf-8')
        texto='# Coleta logistica — resultados preliminares\n\nDistancias OSRM entre centroides ajustados a malha; nao sao percursos entre o CD e clientes. Tempo exclui transito atual, paradas, entregas e intervalos. Sequencias da etapa 22 avaliadas, nao reotimizadas em rodovias. Tarifas nao fornecidas pelo OSRM.\n\n'
        if comb:texto+=f"Gasolina comum Ribeirao Preto: R$ {comb[0]['preco_medio_brl_l']:.2f}/l; {comb[0]['inicio']} a {comb[0]['fim']}; {comb[0]['postos_pesquisados']} postos.\n\n"
        texto+='| Circuito | km rodoviarios preliminares | horas de deslocamento modeladas |\n|---|---:|---:|\n'
        for r in rod:texto+=f"| {r['circuito_id']} | {r['distancia_rodoviaria_preliminar_km']:.2f} | {r['tempo_modelado_deslocamento_h']:.2f} |\n"
        texto+='\nDeslocamentos do centroide ate a malha acima de 500 m foram sinalizados para revisao. Limite e regra de triagem, nao teste de exatidao. Distancias permanecem fora dos parametros finais da etapa 23.\n'
        (out/'RELATORIO.md').write_text(texto,encoding='utf-8')
        conclusao.update(status='COLETA_PARCIAL_COM_PENDENCIAS' if erros else 'COLETA_CONCLUIDA_COM_ROTAS_PRELIMINARES',saida=str(out),bronze=str(raw),erros=erros,contagens={k:len(v) for k,v in tabelas.items()},pedagios_pendentes=True,desgaste_pendente=True,enderecos_pendentes=True,exportacoes=[{'arquivo':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(out.iterdir())])
    except Exception as e:
        conclusao.update(status='FALHA',erro=str(e));raise
    finally:
        conclusao['fim_utc']=datetime.now(timezone.utc).isoformat()
        (q/'conclusao_24.json').write_text(json.dumps(conclusao,ensure_ascii=False,indent=2),encoding='utf-8')
    return conclusao


def main():
    p=argparse.ArgumentParser();p.add_argument('--raiz',type=Path,default=Path(__file__).resolve().parent.parent);p.add_argument('--etapa22',type=Path);p.add_argument('--parametros',type=Path)
    a=p.parse_args();r=executar(a.raiz,datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'_'+uuid4().hex[:8],a.etapa22,a.parametros);print(json.dumps(r,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
