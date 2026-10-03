"""Inventario somente leitura da landing; nao aprova a qualidade dos dados."""
import hashlib
import json
from pathlib import Path


def executar(raiz, destino):
    import openpyxl
    import pdfplumber
    from PIL import Image
    landing = raiz / 'datalake' / '00_landing'
    if not landing.is_dir():
        raise FileNotFoundError(f'Landing inexistente: {landing}')
    arquivos = sorted(p for p in landing.rglob('*') if p.is_file())
    if not arquivos:
        raise ValueError('Landing vazia')
    registros = []
    for p in arquivos:
        r = {'arquivo': p.relative_to(raiz).as_posix(), 'bytes': p.stat().st_size,
             'extensao': p.suffix.lower(), 'status_leitura': 'OK'}
        h = hashlib.sha256()
        with p.open('rb') as f:
            for bloco in iter(lambda: f.read(1024 * 1024), b''):
                h.update(bloco)
        r['sha256'] = h.hexdigest()
        try:
            if r['extensao'] == '.xlsx':
                w = openpyxl.load_workbook(p, read_only=True, data_only=False)
                try:
                    r['abas'] = [{'nome': s.title, 'linhas_ocupadas': sum(any(v is not None for v in row) for row in s.values), 'max_linha': s.max_row, 'max_coluna': s.max_column} for s in w]
                finally:
                    w.close()
            elif r['extensao'] == '.pdf':
                with pdfplumber.open(p) as d:
                    r['paginas'] = [{'pagina': i, 'caracteres_texto': len(s.extract_text() or '')} for i, s in enumerate(d.pages, 1)]
                r['observacao'] = 'Texto extraivel nao garante extracao correta das tabelas; requer conferencia.'
            elif r['extensao'] == '.json':
                dados = json.loads(p.read_text(encoding='utf-8-sig'))
                r['tipo_raiz_json'] = type(dados).__name__
                r['itens_raiz_json'] = len(dados) if isinstance(dados, (list, dict)) else None
                r['observacao'] = 'Sintaxe JSON valida; significado e qualidade dos indicadores avaliados nas etapas especificas.'
            elif r['extensao'] in ('.jpeg', '.jpg', '.png'):
                with Image.open(p) as im:
                    r['largura'], r['altura'] = im.size
                    im.verify()
                r['observacao'] = 'OCR e conferencia visual pendentes.'
            else:
                r['status_leitura'] = 'FORMATO_NAO_SUPORTADO'
        except Exception as e:
            r['status_leitura'] = 'ERRO'
            r['erro'] = str(e)
        registros.append(r)
    resumo = {'etapa': '01_inventario', 'total_arquivos': len(registros),
              'status': 'APROVADO_INVENTARIO' if all(r['status_leitura'] == 'OK' for r in registros) else 'REPROVADO_INVENTARIO',
              'qualidade_dos_dados': 'NAO_AVALIADA', 'arquivos': registros}
    (destino / 'inventario_fontes.json').write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding='utf-8')
    linhas = ['# Inventario das fontes', '', 'Esta etapa verifica leitura e estrutura. Nao valida valores nem fontes.', '', '| Arquivo | Formato | Leitura |', '|---|---|---|']
    linhas += [f"| {r['arquivo']} | {r['extensao']} | {r['status_leitura']} |" for r in registros]
    (destino / 'inventario_fontes.md').write_text('\n'.join(linhas), encoding='utf-8')
    if resumo['status'] != 'APROVADO_INVENTARIO':
        raise RuntimeError('Inventario reprovado; consultar relatorio')
    return resumo
